"""Raspa el catálogo público de series de Ivrea a un Excel, con sus portadas.

    python scripts/demo/raspar_ivrea.py [--limite N]

**Sólo para la demostración en clase** (plan 00, decisión 7). Los datos son de
un catálogo público de editorial, no se publican y se borran al terminar: viven
en `scripts/demo/datos/`, que está fuera del control de versiones.

Ivrea publica una ficha por **serie** (`/titulo/<slug>/`): título, autor,
géneros, demografía, estado, sinopsis y portada. No publica ISBN, precio ni la
lista de tomos; la cantidad de tomos sale del texto «N TOMOS PUBLICADOS» del
estado en Argentina. Lo que falta lo completa `poblar.py`, marcado como sintético.

Una petición por segundo, con un agente identificable, y sólo lo que
`robots.txt` permite. Si una ficha falla, se anota y se sigue: el Excel dice
cuáles quedaron sin leer.
"""

import argparse
import html
import re
import subprocess  # nosec B404 - curl con argumentos fijos
import sys
import time
from pathlib import Path

from openpyxl import Workbook

RAIZ = Path(__file__).resolve().parent
DATOS = RAIZ / "datos"
PORTADAS = DATOS / "portadas"
EXCEL = DATOS / "ivrea.xlsx"
SITEMAP = "https://www.ivrea.com.ar/wp-sitemap-posts-wp-manga-1.xml"
AGENTE = "LongboxDemo/1.0 (trabajo final de grado; una peticion por segundo)"
PAUSA = 1.0

COLUMNAS = [
    "slug",
    "url",
    "titulo",
    "titulo_original",
    "autor",
    "demografia",
    "generos",
    "etiquetas",
    "estado_argentina",
    "tomos_argentina",
    "estado_japon",
    "tomos_japon",
    "formato",
    "sinopsis",
    "portada_url",
    "portada_archivo",
    "error",
]


def pedir(url):
    """Por `curl`, sin seguir redirecciones.

    El sitio contesta «302 a la portada» con el contenido pedido en el cuerpo de
    esa misma respuesta: seguir la redirección trae la portada en su lugar.
    `curl` viene con Windows.
    """
    resultado = subprocess.run(  # nosec B603 B607 - argumentos fijos, sin shell
        ["curl", "-sf", "--max-time", "60", "-A", AGENTE, url],
        capture_output=True,
        check=True,
    )
    return resultado.stdout


def texto(fragmento):
    sin_etiquetas = re.sub(r"<[^>]+>", " ", fragmento)
    return re.sub(r"\s+", " ", html.unescape(sin_etiquetas)).strip()


def resumen(pagina):
    """Los pares «Author(s) => …» del bloque de la serie."""
    pares = {}
    patron = r'summary-heading">\s*<h5>\s*(.*?)\s*</h5>.*?summary-content">(.*?)</div>'
    for clave, valor in re.findall(patron, pagina, re.S):
        pares[texto(clave)] = texto(valor)
    return pares


def descripcion(pagina):
    inicio = pagina.find('class="description-summary"')
    if inicio < 0:
        return ""
    fin = pagina.find('class="c-content-readmore"', inicio)
    bloque = pagina[inicio : fin if fin > 0 else inicio + 20000]
    return texto(bloque.split(">", 1)[1] if ">" in bloque else bloque)


def campo(descripcion_texto, etiqueta):
    """El valor que sigue a «ETIQUETA:» hasta el próximo separador «•»."""
    encontrado = re.search(rf"{etiqueta}\s*:\s*([^•]+)", descripcion_texto, re.I)
    return encontrado.group(1).strip() if encontrado else ""


def tomos(estado):
    if not estado:
        return None
    if re.search(r"TOMO\s+[ÚU]NICO", estado, re.I):
        return 1
    encontrado = re.search(r"(\d+)\s+TOMOS?", estado, re.I)
    return int(encontrado.group(1)) if encontrado else None


def formato(descripcion_texto):
    encontrado = re.search(
        r"\b(B6|A5|A6|B5|Tank[oō]bon|Kanzenban|Big Size|Deluxe)[^.•]{0,60}",
        descripcion_texto,
        re.I,
    )
    return encontrado.group(0).strip() if encontrado else ""


def sinopsis(descripcion_texto):
    """Lo que queda después de los datos de la serie: el texto de presentación."""
    partes = [parte.strip() for parte in descripcion_texto.split("•")]
    ultima = partes[-1] if partes else ""
    # La última parte trae el estado y, a veces, la sinopsis pegada detrás.
    resto = re.sub(r"^.*?TOMOS? PUBLICADOS|^.*?TOMO [ÚU]NICO\)?", "", ultima, flags=re.I)
    return resto.strip()


def leer_serie(url):
    pagina = pedir(url).decode("utf-8", "replace")
    pares = resumen(pagina)
    desc = descripcion(pagina)
    titulo = texto(re.search(r'<div class="post-title">\s*<h1>(.*?)</h1>', pagina, re.S).group(1))
    imagen = re.search(
        r'<div class="summary_image">.*?<img[^>]+?(?:data-src|src)="([^"]+)"', pagina, re.S
    )
    portada = imagen.group(1) if imagen else ""
    # La ficha muestra una miniatura «-193x278»; sin el sufijo está el original.
    portada = re.sub(r"-\d+x\d+(\.\w+)$", r"\1", portada)
    estado_ar = campo(desc, "ESTADO EN ARGENTINA") or campo(desc, "ESTADO")
    estado_jp = campo(desc, "ESTADO EN JAP[ÓO]N")
    return {
        "titulo": titulo,
        "titulo_original": campo(desc, r"NOMBRE ORIGINAL(?: JAPON[ÉE]S)?"),
        "autor": pares.get("Author(s)", "") or campo(desc, "AUTORE?S?"),
        "demografia": campo(desc, "DEMOGRAF[ÍI]A") or pares.get("Type", ""),
        "generos": pares.get("Genre(s)", ""),
        "etiquetas": pares.get("Tag(s)", ""),
        "estado_argentina": estado_ar,
        "tomos_argentina": tomos(estado_ar),
        "estado_japon": estado_jp,
        "tomos_japon": tomos(estado_jp),
        "formato": formato(desc),
        "sinopsis": sinopsis(desc),
        "portada_url": portada,
    }


def bajar_portada(slug, url):
    if not url:
        return ""
    # La extensión sale de la URL sin parámetros, y sólo si es de imagen: algunas
    # traen caracteres que Windows no admite en un nombre de archivo.
    extension = re.search(r"\.(jpe?g|png|webp)(?:$|[?#])", url, re.I)
    slug_seguro = re.sub(r"[^\w.-]", "-", slug)
    destino = PORTADAS / f"{slug_seguro}.{extension.group(1).lower() if extension else 'jpg'}"
    if not destino.exists():
        destino.write_bytes(pedir(url))
        time.sleep(PAUSA)
    return destino.name


def reintentar():
    """Vuelve a leer las filas con error del Excel y las reemplaza en su lugar."""
    from openpyxl import load_workbook

    libro = load_workbook(EXCEL)
    hoja = libro.active
    columnas = [celda.value for celda in hoja[1]]
    arregladas = 0
    for fila in hoja.iter_rows(min_row=2):
        datos = {columna: celda.value for columna, celda in zip(columnas, fila, strict=False)}
        if not datos.get("error"):
            continue
        nuevo = {"slug": datos["slug"], "url": datos["url"], "error": ""}
        try:
            nuevo.update(leer_serie(datos["url"]))
            time.sleep(PAUSA)
            nuevo["portada_archivo"] = bajar_portada(datos["slug"], nuevo["portada_url"])
            arregladas += 1
        except Exception as fallo:  # noqa: BLE001 - se anota y se sigue
            nuevo["error"] = f"{type(fallo).__name__}: {fallo}"[:200]
        for columna, celda in zip(columnas, fila, strict=False):
            celda.value = nuevo.get(columna, "")
    libro.save(EXCEL)
    print(f"Reintentadas: {arregladas} arregladas")
    return 0


def main():
    argumentos = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    argumentos.add_argument("--limite", type=int, default=0, help="Sólo las primeras N series.")
    argumentos.add_argument(
        "--reintentar", action="store_true", help="Volver a leer sólo las series que fallaron."
    )
    opciones = argumentos.parse_args()
    if opciones.reintentar:
        return reintentar()

    PORTADAS.mkdir(parents=True, exist_ok=True)
    mapa = pedir(SITEMAP).decode("utf-8")
    urls = re.findall(r"<loc>([^<]+/titulo/[^<]+)</loc>", mapa)
    if opciones.limite:
        urls = urls[: opciones.limite]
    print(f"Series en el sitemap: {len(urls)}")

    libro = Workbook()
    hoja = libro.active
    hoja.title = "series"
    hoja.append(COLUMNAS)
    errores = 0
    for indice, url in enumerate(urls, 1):
        slug = url.rstrip("/").rsplit("/", 1)[-1]
        fila = {"slug": slug, "url": url}
        try:
            fila.update(leer_serie(url))
            time.sleep(PAUSA)
            fila["portada_archivo"] = bajar_portada(slug, fila["portada_url"])
        except Exception as fallo:  # noqa: BLE001 - se anota y se sigue con la próxima
            fila["error"] = f"{type(fallo).__name__}: {fallo}"[:200]
            errores += 1
        hoja.append([fila.get(columna, "") for columna in COLUMNAS])
        if indice % 25 == 0:
            libro.save(EXCEL)
            print(f"  {indice}/{len(urls)} ({errores} con error)", flush=True)
    libro.save(EXCEL)
    print(f"Listo: {EXCEL} · {len(urls)} series, {errores} con error")


if __name__ == "__main__":
    sys.exit(main())
