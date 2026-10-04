"""Correos de novedades de editorial, armados con los datos de Ivrea. Sólo desarrollo.

    python scripts/demo/enviar_novedades.py --cantidad 5 --carpeta scripts/demo/salida/correos
    python scripts/demo/enviar_novedades.py --cantidad 1 --smtp localhost:1025
        --para casilla@ejemplo.com --remitente novedades@ivrea.demo

Arma correos con la forma de los boletines de novedades de una editorial: una
tabla con un producto por fila (título, tomo, formato, fecha de salida, precio)
y la versión en texto. Los deja como `.eml`, que el sistema ya ingiere con
`manage.py ingerir_correos <carpeta>`, o los manda por SMTP a una casilla.

Es la herramienta de la demostración (plan 00, decisión 8): **no es parte del
sistema** y no lo toca. La lectura de la casilla por IMAP es `CO-016`, que no
existe todavía.
"""

import argparse
import random
import smtplib
import sys
from datetime import date, datetime, time, timedelta
from email.message import EmailMessage
from email.utils import format_datetime, make_msgid
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent
EXCEL = RAIZ / "datos" / "ivrea.xlsx"


def series(excel=EXCEL):
    """Las series del Excel, sin las que fallaron al leerse."""
    hoja = load_workbook(excel, read_only=True).active
    filas = hoja.iter_rows(values_only=True)
    columnas = next(filas)
    salida = []
    for fila in filas:
        registro = dict(zip(columnas, fila, strict=False))
        if registro.get("titulo") and not registro.get("error"):
            salida.append(registro)
    return salida


def _pesos(valor):
    return f"$ {valor:,.0f}".replace(",", ".")


def armar(remitente, destinatario, renglones, fecha_envio, *, asunto=None, message_id=None):
    """Un correo de novedades. `renglones` son dicts con titulo, tomo, formato, salida y precio."""
    correo = EmailMessage()
    correo["From"] = remitente
    correo["To"] = destinatario
    correo["Subject"] = asunto or f"Novedades del {fecha_envio:%d/%m/%Y}"
    correo["Date"] = format_datetime(datetime.combine(fecha_envio, time(9, 0)).astimezone())
    correo["Message-ID"] = message_id or make_msgid(domain="novedades.demo")

    texto = [f"Novedades para la semana del {fecha_envio:%d/%m/%Y}.", ""]
    for renglon in renglones:
        texto.append(
            f"- {renglon['titulo']} #{renglon['tomo']} · {renglon['formato']} · "
            f"sale el {renglon['salida']:%d/%m/%Y} · {_pesos(renglon['precio'])}"
        )
    texto += ["", "Saludos, el equipo comercial."]
    correo.set_content("\n".join(texto))

    filas = "".join(
        f"<tr><td>{r['titulo']}</td><td>{r['tomo']}</td><td>{r['formato']}</td>"
        f"<td>{r['salida']:%d/%m/%Y}</td><td>{_pesos(r['precio'])}</td></tr>"
        for r in renglones
    )
    correo.add_alternative(
        "<p>Novedades para la semana del "
        f"{fecha_envio:%d/%m/%Y}.</p>"
        "<table border='1'><tr><th>Título</th><th>Tomo</th><th>Formato</th>"
        f"<th>Salida</th><th>Precio</th></tr>{filas}</table>"
        "<p>Saludos, el equipo comercial.</p>",
        subtype="html",
    )
    return correo


def renglones_al_azar(azar, catalogo, cantidad, desde):
    elegidas = azar.sample(catalogo, min(cantidad, len(catalogo)))
    return [
        {
            "titulo": serie["titulo"],
            "tomo": (serie.get("tomos_argentina") or 0) + 1,
            "formato": "B6" if "B6" in (serie.get("formato") or "") else "Tankoubon",
            "salida": desde + timedelta(days=azar.randint(3, 40)),
            "precio": azar.choice([14900, 16500, 18500, 21900, 24900]),
        }
        for serie in elegidas
    ]


def main():
    argumentos = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    argumentos.add_argument("--cantidad", type=int, default=1)
    argumentos.add_argument("--renglones", type=int, default=6)
    argumentos.add_argument("--remitente", default="novedades@ivrea.demo")
    argumentos.add_argument("--para", default="casilla@longbox.demo")
    argumentos.add_argument("--carpeta", type=Path, help="Dejar los .eml en esta carpeta.")
    argumentos.add_argument("--smtp", help="Servidor SMTP host:puerto para mandarlos de verdad.")
    argumentos.add_argument("--semilla", type=int, default=20261005)
    opciones = argumentos.parse_args()

    azar = random.Random(opciones.semilla)  # nosec B311 - datos de muestra, no criptografía
    catalogo = series()
    hoy = date.today()
    correos = [
        armar(
            opciones.remitente,
            opciones.para,
            renglones_al_azar(azar, catalogo, opciones.renglones, hoy),
            hoy,
        )
        for _ in range(opciones.cantidad)
    ]
    if opciones.carpeta:
        opciones.carpeta.mkdir(parents=True, exist_ok=True)
        for indice, correo in enumerate(correos, 1):
            (opciones.carpeta / f"novedades-{indice:03d}.eml").write_bytes(bytes(correo))
        print(f"{len(correos)} correo(s) en {opciones.carpeta}")
    if opciones.smtp:
        servidor, _, puerto = opciones.smtp.partition(":")
        with smtplib.SMTP(servidor, int(puerto or 25), timeout=20) as conexion:
            for correo in correos:
                conexion.send_message(correo)
        print(f"{len(correos)} correo(s) enviados por {opciones.smtp}")
    if not (opciones.carpeta or opciones.smtp):
        print(bytes(correos[0]).decode("utf-8", "replace"))


if __name__ == "__main__":
    sys.exit(main())
